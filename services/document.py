"""DOCX template handling with python-docx."""
from __future__ import annotations

from pathlib import Path

from docx import Document

PLACEHOLDER_1 = "[COMPANY_PARAGRAPH_1]"
PLACEHOLDER_2 = "[COMPANY_PARAGRAPH_2]"
PLACEHOLDERS = (PLACEHOLDER_1, PLACEHOLDER_2)

# Validation guardrails (MVP): reject empty output and absurdly long paragraphs.
MIN_CHARS = 80
MAX_CHARS = 1800


def _iter_paragraphs(doc: Document):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def template_has_placeholders(template_path: str | Path) -> bool:
    doc = Document(str(template_path))
    full = "\n".join(p.text for p in _iter_paragraphs(doc))
    return all(ph in full for ph in PLACEHOLDERS)


def _set_paragraph_text(paragraph, new_text: str) -> None:
    """Replace paragraph content while keeping the first run's style when possible."""
    if paragraph.runs:
        style = paragraph.runs[0].style
        paragraph.text = new_text
        try:
            paragraph.runs[0].style = style
        except Exception:
            pass
    else:
        paragraph.text = new_text


def replace_placeholders(
    template_path: str | Path,
    output_path: str | Path,
    paragraphs: dict[str, str],
) -> None:
    mapping = {
        PLACEHOLDER_1: paragraphs["company_paragraph_1"],
        PLACEHOLDER_2: paragraphs["company_paragraph_2"],
    }
    doc = Document(str(template_path))
    replaced = {ph: False for ph in PLACEHOLDERS}

    for paragraph in _iter_paragraphs(doc):
        for placeholder, new_text in mapping.items():
            if placeholder in paragraph.text:
                # Full-paragraph placeholder (normal case): clean replacement.
                if paragraph.text.strip() == placeholder:
                    _set_paragraph_text(paragraph, new_text)
                else:
                    # Inline placeholder: string-level replacement, may split runs
                    # but preserves surrounding fixed content.
                    for run in paragraph.runs:
                        if placeholder in run.text:
                            run.text = run.text.replace(placeholder, new_text)
                    if placeholder in paragraph.text:  # placeholder spanned runs
                        paragraph.text = paragraph.text.replace(placeholder, new_text)
                replaced[placeholder] = True

    remaining = [ph for ph, done in replaced.items() if not done]
    if remaining:
        raise RuntimeError(f"Placeholders not found in template: {remaining}")

    # Safety net: no placeholder may survive in the final document.
    full = "\n".join(p.text for p in _iter_paragraphs(doc))
    leftover = [ph for ph in PLACEHOLDERS if ph in full]
    if leftover:
        raise RuntimeError(f"Placeholders remain after replacement: {leftover}")

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))


def validate_paragraphs(paragraphs: dict[str, str]) -> None:
    """MVP validation before rendering. Raises RuntimeError with a clear message."""
    p1 = (paragraphs.get("company_paragraph_1") or "").strip()
    p2 = (paragraphs.get("company_paragraph_2") or "").strip()
    if not p1 or not p2:
        raise RuntimeError(
            "Validation failed: both company_paragraph_1 and company_paragraph_2 "
            "must exist and be non-empty."
        )
    for name, text in (("company_paragraph_1", p1), ("company_paragraph_2", p2)):
        if any(ph in text for ph in PLACEHOLDERS):
            raise RuntimeError(f"Validation failed: {name} still contains a placeholder.")
        if len(text) < MIN_CHARS:
            raise RuntimeError(
                f"Validation failed: {name} is too short ({len(text)} chars, "
                f"minimum {MIN_CHARS})."
            )
        if len(text) > MAX_CHARS:
            raise RuntimeError(
                f"Validation failed: {name} is too long ({len(text)} chars, "
                f"maximum {MAX_CHARS}). Regenerate with a shorter text."
            )
